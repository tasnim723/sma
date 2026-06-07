import base64
import io
from pdfminer.high_level import extract_text as extract_pdf_text
from pdfminer.high_level import extract_pages
from pdfminer.layout import LTTextContainer

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from PIL import Image
except ImportError:
    Image = None


def is_pdf_bytes(data: bytes) -> bool:
    """Check PDF magic bytes (%PDF-)"""
    return data[:5] == b'%PDF-'


from pdfminer.pdfparser import PDFParser
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdftypes import resolve1

def extract_pdf_links(pdf_bytes: bytes) -> str:
    links = []
    try:
        parser = PDFParser(io.BytesIO(pdf_bytes))
        doc = PDFDocument(parser)
        if 'Pages' in doc.catalog:
            pages = resolve1(doc.catalog['Pages'])
            if 'Kids' in pages:
                kids = resolve1(pages['Kids'])
                for page_ref in kids:
                    page = resolve1(page_ref)
                    if 'Annots' in page:
                        annots = resolve1(page['Annots'])
                        # Handle case where annots is not a list
                        if not isinstance(annots, list):
                            annots = [annots]
                        for annot_ref in annots:
                            annot = resolve1(annot_ref)
                            if isinstance(annot, dict) and 'A' in annot:
                                action = resolve1(annot['A'])
                                if isinstance(action, dict) and 'URI' in action:
                                    uri = action['URI']
                                    if isinstance(uri, bytes):
                                        links.append(uri.decode('utf-8', 'ignore'))
                                    else:
                                        links.append(str(uri))
    except Exception as e:
        pass
    
    if links:
        important_links = [l for l in links if 'linkedin.com' in l.lower() or 'github.com' in l.lower()]
        if important_links:
            return "\n\n--- EXTRACTED HYPERLINKS ---\n" + "\n".join(important_links)
    return ""


def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
    """Extract all text from PDF bytes using pdfminer."""
    try:
        file_io = io.BytesIO(pdf_bytes)
        text = extract_pdf_text(file_io)
        extracted_links = extract_pdf_links(pdf_bytes)
        if text and text.strip():
            return text.strip() + extracted_links
        # Fallback: try page-by-page extraction
        file_io.seek(0)
        pages_text = []
        for page_layout in extract_pages(file_io):
            for element in page_layout:
                if isinstance(element, LTTextContainer):
                    pages_text.append(element.get_text())
        return '\n'.join(pages_text).strip() + extracted_links
    except Exception as e:
        return f"[PDF_ERROR: {str(e)}]"


def process_file_data(filename: str, file_type: str, base64_data: str) -> str:
    """Extract text from base64-encoded or raw files."""
    try:
        # Strip data-URL header if present
        if "," in base64_data:
            _, base64_data = base64_data.split(",", 1)

        file_bytes = base64.b64decode(base64_data)

        # Detect PDF by magic bytes (more reliable than MIME type)
        if is_pdf_bytes(file_bytes):
            return extract_text_from_pdf_bytes(file_bytes)

        if file_type == "application/pdf" or filename.lower().endswith(".pdf"):
            return extract_text_from_pdf_bytes(file_bytes)

        if ("word" in file_type or filename.lower().endswith(".docx")) and Document:
            doc = Document(io.BytesIO(file_bytes))
            return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])

        if "text" in file_type or filename.lower().endswith(".txt"):
            return file_bytes.decode("utf-8", errors="ignore")

        if "image" in file_type and Image:
            img = Image.open(io.BytesIO(file_bytes))
            return f"[IMAGE: {filename} ({img.format}, {img.size[0]}x{img.size[1]})]"

        return f"[UNSUPPORTED FILE: {filename} ({file_type})]"

    except Exception as e:
        return f"[ERROR PROCESSING {filename}: {str(e)}]"
